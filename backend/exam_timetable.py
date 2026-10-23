"""
UniBot — Exam timetable.

Built for the DKUT CS exam timetable format:
  - Title row: "…EXAMINATION TIMETABLE AUGUST 2026" (or SUPPLEMENTARY …)
  - Header row: DAY | DATE | YEAR | CODE | UNIT | TIME | VENUE | PR. INVIGILATOR | ASS. INVIGILATOR
  - One sheet for main exams, one for supplementary exams
  - Merged cells: one venue for many rows, one day/date/year for a morning
    and an afternoon paper

Only what students need is kept: unit, date, time, venue.
Invigilator names are never read or stored.

Each upload replaces the whole exam timetable. When a new version is
published, the changes (moved, added, cancelled papers) are worked out
so the affected year groups can be told exactly what changed.

Validated against: CS_Exam_Timetable_August_2026_Draft_1
"""

import re
import logging
from datetime import datetime, date

import openpyxl

from database import get_db, _ph

logger = logging.getLogger("unibot.exams")

UNIT_CODE_RE = re.compile(r"\b([A-Z]{3})\s*(\d{4})\b")
VERSION_RE = re.compile(r"\b(DRAFT\s*\d+|FINAL|REVISED(?:\s*\d+)?)\b", re.I)
PERIOD_RE = re.compile(r"(?:EXAMINATION|EXAM)S?\s+TIMETABLE\s+(.+)$", re.I)
WEEKDAYS = ["MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY", "SATURDAY", "SUNDAY"]
NEEDED = ["DAY", "DATE", "YEAR", "CODE", "UNIT", "TIME", "VENUE"]


# ─────────────────────────────────────────────────────────────
# Reading the file
# ─────────────────────────────────────────────────────────────
def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return re.sub(r"\s+", " ", str(v)).strip()


def _merged_value_map(ws):
    """Every cell inside a merged block gets the value of the block's top-left cell."""
    filled = {}
    for rng in ws.merged_cells.ranges:
        top_left = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                filled[(r, c)] = top_left
    return filled


def _parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _text(v)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def normalise_time(s: str) -> str:
    """'11:00AM-1:00PM', '8:00AM - 6:00PM' -> '11:00AM-1:00PM' style."""
    s = re.sub(r"\s*-\s*", "-", s.upper().replace("–", "-"))
    return re.sub(r"\s+", "", s)


def time_start_minutes(slot: str) -> int:
    m = re.match(r"(\d{1,2})(?::(\d{2}))?(AM|PM)", slot or "")
    if not m:
        return 24 * 60
    h, mins, half = int(m.group(1)) % 12, int(m.group(2) or 0), m.group(3)
    return (h + (12 if half == "PM" else 0)) * 60 + mins


def parse_exam_excel(file_path: str, filename: str = "") -> dict:
    """
    Returns {exams: [...], label: str, errors: [...], warnings: [...]}
    Each exam: kind (MAIN/SUPP), exam_date (YYYY-MM-DD), day, year_group,
               unit_code, unit_name, time_slot, venue
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    exams, errors, warnings = [], [], []
    period = ""

    for ws in wb.worksheets:
        merged = _merged_value_map(ws)

        def val(r, c):
            return merged.get((r, c), ws.cell(r, c).value)

        # Find the header row (the one that says DAY … DATE … CODE)
        header_row, cols = None, {}
        for r in range(1, min(ws.max_row, 15) + 1):
            names = {_text(ws.cell(r, c).value).upper(): c for c in range(1, ws.max_column + 1)}
            if {"DAY", "DATE", "CODE"} <= set(names):
                header_row = r
                cols = {k: names[k] for k in NEEDED if k in names}
                break
        if not header_row:
            continue

        title = " ".join(_text(ws.cell(r, 1).value) for r in range(1, header_row))
        kind = "SUPP" if re.search(r"SUPP", title + " " + ws.title, re.I) else "MAIN"
        if not period:
            pm = PERIOD_RE.search(title)
            if pm:
                period = pm.group(1).strip()

        for r in range(header_row + 1, ws.max_row + 1):
            code_raw = _text(val(r, cols.get("CODE", 0))) if "CODE" in cols else ""
            code_m = UNIT_CODE_RE.search(code_raw.upper())
            if not code_m:
                continue                    # empty slot, title or footer row
            code = f"{code_m.group(1)} {code_m.group(2)}"

            d = _parse_date(val(r, cols["DATE"]))
            day = _text(val(r, cols.get("DAY", 0))).upper() if "DAY" in cols else ""
            year_m = re.search(r"(\d)", _text(val(r, cols.get("YEAR", 0)))) if "YEAR" in cols else None
            time_slot = normalise_time(_text(val(r, cols.get("TIME", 0)))) if "TIME" in cols else ""
            venue = _text(val(r, cols.get("VENUE", 0))).upper() if "VENUE" in cols else ""
            name = _text(val(r, cols.get("UNIT", 0))) if "UNIT" in cols else ""
            where = f"{ws.title} row {r} ({code})"

            if not d:
                errors.append(f"{where}: date could not be read")
                continue
            if not year_m:
                errors.append(f"{where}: year could not be read")
                continue
            # Garbage-in check: the DAY column must match the DATE
            real_day = WEEKDAYS[d.weekday()]
            if day and day != real_day:
                errors.append(f"{where}: says {day.title()} but {d:%d %b %Y} is a {real_day.title()}")
            if not time_slot:
                warnings.append(f"{where}: no time given")
            if not venue:
                warnings.append(f"{where}: no venue given")

            exams.append({
                "kind": kind, "exam_date": d.isoformat(), "day": real_day,
                "year_group": int(year_m.group(1)), "unit_code": code, "unit_name": name,
                "time_slot": time_slot, "venue": venue or "TBA",
            })

    version_m = VERSION_RE.search(re.sub(r"[_\-.]+", " ", filename or ""))
    version = re.sub(r"\s+", " ", version_m.group(1).upper()) if version_m else ""
    version = re.sub(r"^(DRAFT|REVISED)\s*(\d+)$", r"\1 \2", version)
    label = " ".join(p for p in [period.upper(), version] if p) or "EXAM TIMETABLE"
    if not exams:
        errors.insert(0, "No exams found. Is this the exam timetable "
                         "(DAY | DATE | YEAR | CODE | UNIT | TIME | VENUE)?")
    return {"exams": exams, "label": label, "errors": errors, "warnings": warnings}


# ─────────────────────────────────────────────────────────────
# Checks before publishing
# ─────────────────────────────────────────────────────────────
def exam_clashes(exams):
    """Same year, two main papers at the same date and time."""
    seen, clashes = {}, []
    for e in exams:
        if e["kind"] != "MAIN":
            continue
        key = (e["year_group"], e["exam_date"], e["time_slot"])
        if key in seen and seen[key]["unit_code"] != e["unit_code"]:
            o = seen[key]
            clashes.append({
                "type": "STUDENT",
                "detail": f"Year {e['year_group']} on {e['exam_date']} {e['time_slot']}: "
                          f"{o['unit_code']} and {e['unit_code']} at the same time"})
        seen.setdefault(key, e)
    return clashes


def describe(e) -> str:
    d = datetime.strptime(e["exam_date"], "%Y-%m-%d")
    return f"{d:%a %d %b}, {e['time_slot']}, {e['venue']}"


def diff_exams(old, new):
    """
    What changed between two versions, per year group (main exams only):
    {year: [change lines]}
    """
    def index(rows):
        return {(e["year_group"], e["unit_code"]): e for e in rows if e["kind"] == "MAIN"}

    o, n = index(old), index(new)
    changes = {}
    for key, e in n.items():
        before = o.get(key)
        if before is None:
            line = f"➕ {e['unit_code']} {e['unit_name']}: {describe(e)}"
        elif (before["exam_date"], before["time_slot"], before["venue"]) != \
                (e["exam_date"], e["time_slot"], e["venue"]):
            line = (f"🔁 {e['unit_code']} {e['unit_name']}\n"
                    f"   was {describe(before)}\n   now {describe(e)}")
        else:
            continue
        changes.setdefault(key[0], []).append(line)
    for key, e in o.items():
        if key not in n:
            changes.setdefault(key[0], []).append(
                f"❌ {e['unit_code']} {e['unit_name']} is no longer on the timetable")
    return changes


# ─────────────────────────────────────────────────────────────
# Database
# ─────────────────────────────────────────────────────────────
def _f(r, name, idx):
    return r[name] if hasattr(r, "keys") else r[idx]


EXAM_COLS = ["kind", "exam_date", "day", "year_group", "unit_code", "unit_name",
             "time_slot", "venue"]


def get_all_exams():
    with get_db() as conn:
        cur = conn.cursor()
        cur.execute(f"SELECT {', '.join(EXAM_COLS)} FROM exams")
        rows = cur.fetchall()
    out = []
    for r in rows:
        e = {c: _f(r, c, i) for i, c in enumerate(EXAM_COLS)}
        e["exam_date"] = str(e["exam_date"])[:10]
        e["year_group"] = int(e["year_group"])
        out.append(e)
    return out


def save_exams(parsed: dict, filename: str, official_link: str, uploaded_by: str):
    """Replace the exam timetable. Returns the changes per year compared with the old one."""
    old = get_all_exams()
    changes = diff_exams(old, parsed["exams"]) if old else {}
    if not official_link:                       # keep the last official link if none given
        previous = get_exam_info()
        official_link = (previous or {}).get("official_link") or ""
    with get_db() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM exams")
        for e in parsed["exams"]:
            cur.execute(_ph(f"""INSERT INTO exams ({', '.join(EXAM_COLS)})
                            VALUES ({', '.join(['%s'] * len(EXAM_COLS))})"""),
                        tuple(e[c] for c in EXAM_COLS))
        cur.execute(_ph("""INSERT INTO exam_uploads (label, filename, official_link, exams, uploaded_by)
                       VALUES (%s, %s, %s, %s, %s)"""),
                    (parsed["label"], filename, official_link, len(parsed["exams"]), uploaded_by))
    logger.info(f"Exam timetable published: {parsed['label']} ({len(parsed['exams'])} papers)")
    return {"first_upload": not old, "changes": changes}


def get_exam_info():
    """The live exam timetable's label, link and upload time (or None)."""
    with get_db() as conn:
        cur = conn.cursor()
        cur.execute("""SELECT label, filename, official_link, exams, uploaded_by, uploaded_at
                       FROM exam_uploads ORDER BY id DESC LIMIT 1""")
        r = cur.fetchone()
    if not r:
        return None
    keys = ["label", "filename", "official_link", "exams", "uploaded_by", "uploaded_at"]
    return {k: _f(r, k, i) for i, k in enumerate(keys)}


def exams_for_year(year: int, include_supp: bool = False):
    """A year's exams in date and time order."""
    rows = [e for e in get_all_exams()
            if e["year_group"] == year and (include_supp or e["kind"] == "MAIN")]
    return sorted(rows, key=lambda e: (e["exam_date"], time_start_minutes(e["time_slot"])))